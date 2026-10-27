"""Transparent, source-linked RFP requirement CANDIDATES, not legal conclusions.

No external LLM or network call: text-only PDF and DOCX are supported.
The reviewer must confirm every candidate and its relevance.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass

import fitz
from docx import Document as WordDocument
from fastapi import HTTPException

MAX_DOCUMENT_BYTES = 8 * 1024 * 1024
MAX_EXTRACTED_CHARS = 1_000_000
MAX_REQUIREMENTS = 250
KEYWORDS = re.compile(
    r"\b(musi|muszą|należy|wymaga(?:ne|ny|ją)?|obowiązkow[ey]|"
    r"warunek|kryterium|termin|doświadczenie|zaświadczenie|"
    r"shall|must|required|mandatory|deadline|qualification|"
    r"provide|submit|evidence|certificate|compliance)\b",
    re.IGNORECASE,
)
DEADLINE = re.compile(r"\b(termin|deadline|do\s+dnia|submit\s+by)\b", re.I)
LEGAL = re.compile(r"\b(ubezpieczen|certyfikat|oświadczen|licencj|"
                   r"certificate|insurance|licen[cs]e|declaration)\w*", re.I)
COMMERCIAL = re.compile(r"\b(cen[ayę]|wycen|ofert[ayę]|price|pricing|budget|"
                        r"commercial|płatnoś)\w*", re.I)


@dataclass(frozen=True)
class Candidate:
    locator: str
    evidence: str
    category: str


def category(text: str) -> str:
    if DEADLINE.search(text):
        return "deadline"
    if LEGAL.search(text):
        return "legal"
    if COMMERCIAL.search(text):
        return "commercial"
    return "technical"


def candidate_segments(filename: str, content: bytes) -> list[Candidate]:
    if len(content) > MAX_DOCUMENT_BYTES or not content:
        raise HTTPException(413, "Document must be 1 byte to 8 MiB")
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    segments: list[tuple[str, str]] = []
    if suffix == "pdf":
        try:
            with fitz.open(stream=content, filetype="pdf") as document:
                if document.page_count > 80:
                    raise HTTPException(422, "Maximum 80 PDF pages per upload")
                for index, page in enumerate(document):
                    text = page.get_text("text", sort=True)
                    for line in text.splitlines():
                        segments.append((f"page {index + 1}", line))
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, "Cannot read PDF (scans need OCR)") from exc
    elif suffix == "docx":
        try:
            doc = WordDocument(io.BytesIO(content))
            if len(doc.paragraphs) > 3000 or len(doc.tables) > 100:
                raise HTTPException(422, "DOCX is too long")
            segments.extend(
                (f"paragraph {i + 1}", p.text) for i, p in enumerate(doc.paragraphs)
            )
            for table_index, table in enumerate(doc.tables):
                if len(table.rows) > 300:
                    raise HTTPException(422, "DOCX table too large")
                for row_index, row in enumerate(table.rows):
                    cells = [cell.text for cell in row.cells[:12]]
                    segments.append((f"table {table_index+1}, row {row_index+1}", " | ".join(cells)))
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, "Cannot read DOCX") from exc
    else:
        raise HTTPException(415, "Only text PDFs and DOCX documents are supported")
    result: list[Candidate] = []
    char_count = 0
    for locator, raw in segments:
        char_count += len(raw)
        if char_count > MAX_EXTRACTED_CHARS:
            raise HTTPException(422, "Document text exceeds supported limit")
        normalized = " ".join(raw.split())
        if len(normalized) >= 12 and KEYWORDS.search(normalized):
            result.append(Candidate(locator, normalized[:700], category(normalized)))
            if len(result) >= MAX_REQUIREMENTS:
                break
    return result
