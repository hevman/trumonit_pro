"""Manual, reviewable bid responses. No AI, LLM or auto-generated copy.

Library entries are explicitly approved by an organization owner. Applying one
copies its content into a draft; the resulting response needs its own approval.
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

from docx import Document as WordDocument
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel, Field, model_validator
from sqlalchemy import select
from sqlalchemy.orm import Session

from .db import get_db, now
from .models import AnswerLibraryEntry, RequirementResponse, Requirement, Project
from .projects import owned_project, record
from .security import Principal, principal, require_editor

router = APIRouter(tags=["Manual RFP responses"])


class LibraryDraft(BaseModel):
    title: str = Field(min_length=3, max_length=150)
    category: str = Field(default="general", min_length=1, max_length=40)
    body: str = Field(min_length=15, max_length=10000)
    evidence_reference: str = Field(default="", max_length=400)
    valid_until: date | None = None


class ResponseDraft(BaseModel):
    body: str | None = Field(default=None, min_length=10, max_length=10000)
    library_entry_id: str | None = None

    @model_validator(mode="after")
    def exactly_one_source(self):
        if (self.body is None) == (self.library_entry_id is None):
            raise ValueError("Provide either a manual answer or a library entry ID")
        return self


def require_owner(user: Principal):
    if user.role != "owner":
        raise HTTPException(403, "Only organization owners can approve answers")


def owned_library_entry(db: Session, user: Principal, entry_id: str):
    entry = db.scalar(select(AnswerLibraryEntry).where(
        AnswerLibraryEntry.id == entry_id,
        AnswerLibraryEntry.workspace_id == user.workspace_id,
    ))
    if entry is None:
        raise HTTPException(404, "Answer library entry not found")
    return entry


def owned_requirement(db: Session, user: Principal, project_id: str, requirement_id: str):
    owned_project(db, project_id, user)
    item = db.scalar(select(Requirement).where(
        Requirement.id == requirement_id,
        Requirement.project_id == project_id,
        Requirement.workspace_id == user.workspace_id,
    ))
    if item is None:
        raise HTTPException(404, "Requirement not found")
    return item


def entry_data(entry: AnswerLibraryEntry):
    valid = entry.valid_until is None or date.fromisoformat(entry.valid_until) >= date.today()
    return {
        "id": entry.id, "title": entry.title, "category": entry.category,
        "body": entry.body, "evidence_reference": entry.evidence_reference,
        "valid_until": entry.valid_until, "status": entry.status,
        "usable": entry.status == "approved" and valid,
        "approved_by": entry.approved_by,
        "approved_at": entry.approved_at.isoformat() if entry.approved_at else None,
    }


def response_data(row: RequirementResponse):
    return {
        "id": row.id, "requirement_id": row.requirement_id,
        "body": row.body, "status": row.status,
        "library_entry_id": row.library_entry_id,
        "approved_by": row.approved_by,
        "approved_at": row.approved_at.isoformat() if row.approved_at else None,
    }


@router.get("/library")
def list_library(user: Principal = Depends(principal), db: Session = Depends(get_db)):
    entries = db.scalars(select(AnswerLibraryEntry).where(
        AnswerLibraryEntry.workspace_id == user.workspace_id,
    ).order_by(AnswerLibraryEntry.created_at.desc()).limit(200)).all()
    return [entry_data(entry) for entry in entries]


@router.post("/library", status_code=201)
def create_entry(payload: LibraryDraft,
                 user: Principal = Depends(principal),
                 db: Session = Depends(get_db)):
    require_editor(user)
    entry = AnswerLibraryEntry(
        workspace_id=user.workspace_id, title=payload.title.strip(),
        category=payload.category.strip(), body=payload.body.strip(),
        evidence_reference=payload.evidence_reference.strip(),
        valid_until=payload.valid_until.isoformat() if payload.valid_until else None,
        status="draft",
    )
    db.add(entry)
    db.flush()
    record(db, user, "library.created", entry.id)
    db.commit()
    return entry_data(entry)


@router.patch("/library/{entry_id}")
def edit_entry(entry_id: str, payload: LibraryDraft,
               user: Principal = Depends(principal),
               db: Session = Depends(get_db)):
    require_editor(user)
    entry = owned_library_entry(db, user, entry_id)
    entry.title = payload.title.strip()
    entry.category = payload.category.strip()
    entry.body = payload.body.strip()
    entry.evidence_reference = payload.evidence_reference.strip()
    entry.valid_until = payload.valid_until.isoformat() if payload.valid_until else None
    entry.status = "draft"  # Material edit always requires renewed approval.
    entry.approved_by = None
    entry.approved_at = None
    entry.updated_at = now()
    record(db, user, "library.edited", entry.id)
    db.commit()
    return entry_data(entry)


@router.post("/library/{entry_id}/approve")
def approve_entry(entry_id: str,
                  user: Principal = Depends(principal),
                  db: Session = Depends(get_db)):
    require_owner(user)
    entry = owned_library_entry(db, user, entry_id)
    if entry.valid_until and date.fromisoformat(entry.valid_until) < date.today():
        raise HTTPException(409, "Expired library entry cannot be approved")
    entry.status = "approved"
    entry.approved_by = user.user_id
    entry.approved_at = now()
    entry.updated_at = now()
    record(db, user, "library.approved", entry.id)
    db.commit()
    return entry_data(entry)


@router.get("/projects/{project_id}/responses")
def list_responses(project_id: str,
                   user: Principal = Depends(principal),
                   db: Session = Depends(get_db)):
    owned_project(db, project_id, user)
    rows = db.scalars(select(RequirementResponse).where(
        RequirementResponse.project_id == project_id,
        RequirementResponse.workspace_id == user.workspace_id,
    ).limit(500)).all()
    return [response_data(row) for row in rows]


@router.patch("/projects/{project_id}/requirements/{requirement_id}/response")
def set_response(project_id: str, requirement_id: str, payload: ResponseDraft,
                 user: Principal = Depends(principal),
                 db: Session = Depends(get_db)):
    require_editor(user)
    requirement = owned_requirement(db, user, project_id, requirement_id)
    source_id = None
    if payload.library_entry_id is not None:
        source = owned_library_entry(db, user, payload.library_entry_id)
        if not entry_data(source)["usable"]:
            raise HTTPException(409, "Library answer must be approved and unexpired")
        source_id = source.id
        text = source.body
    else:
        text = (payload.body or "").strip()
        if len(text) < 10:
            raise HTTPException(422, "Answer must contain at least 10 characters")
    row = db.scalar(select(RequirementResponse).where(
        RequirementResponse.requirement_id == requirement.id,
        RequirementResponse.workspace_id == user.workspace_id,
        RequirementResponse.project_id == project_id,
    ))
    if row is None:
        row = RequirementResponse(
            workspace_id=user.workspace_id,
            project_id=project_id,
            requirement_id=requirement.id,
            body=text,
            library_entry_id=source_id,
        )
        db.add(row)
    else:
        row.body = text
        row.library_entry_id = source_id
    row.status = "draft"
    row.approved_by = None
    row.approved_at = None
    row.updated_at = now()
    db.flush()
    record(db, user, "response.drafted", row.id)
    db.commit()
    return response_data(row)


@router.post("/projects/{project_id}/requirements/{requirement_id}/response/approve")
def approve_response(project_id: str, requirement_id: str,
                     user: Principal = Depends(principal),
                     db: Session = Depends(get_db)):
    require_owner(user)
    requirement = owned_requirement(db, user, project_id, requirement_id)
    if requirement.review_status != "met":
        raise HTTPException(409, "Mark this requirement as Met before approving an answer")
    row = db.scalar(select(RequirementResponse).where(
        RequirementResponse.requirement_id == requirement.id,
        RequirementResponse.project_id == project_id,
        RequirementResponse.workspace_id == user.workspace_id,
    ))
    if row is None:
        raise HTTPException(404, "Draft response not found")
    row.status = "approved"
    row.approved_by = user.user_id
    row.approved_at = now()
    row.updated_at = now()
    record(db, user, "response.approved", row.id)
    db.commit()
    return response_data(row)


@router.get("/projects/{project_id}/response-pack")
def export_response_pack(project_id: str,
                         user: Principal = Depends(principal),
                         db: Session = Depends(get_db)):
    project = owned_project(db, project_id, user)
    ready = db.execute(
        select(Requirement, RequirementResponse)
        .join(RequirementResponse, RequirementResponse.requirement_id == Requirement.id)
        .where(
            Requirement.project_id == project.id,
            Requirement.workspace_id == user.workspace_id,
            RequirementResponse.project_id == project.id,
            RequirementResponse.workspace_id == user.workspace_id,
            RequirementResponse.status == "approved",
            Requirement.review_status == "met",
        ).order_by(Requirement.created_at, Requirement.id).limit(500)
    ).all()
    if not ready:
        raise HTTPException(409, "No approved, Met requirement responses available to export")
    for requirement, answer in ready:
        if answer.library_entry_id:
            library = db.scalar(select(AnswerLibraryEntry).where(
                AnswerLibraryEntry.id == answer.library_entry_id,
                AnswerLibraryEntry.workspace_id == user.workspace_id,
            ))
            if library is None or (library.valid_until and date.fromisoformat(library.valid_until) < date.today()):
                raise HTTPException(409, "Review an answer linked to an expired or missing company document before export")
    document = WordDocument()
    document.add_heading("Trumonit | Response working pack", 0)
    document.add_paragraph("INTERNAL REVIEW COPY — Not a submission-ready legal declaration.")
    document.add_heading(project.name, level=1)
    document.add_paragraph(f"Buyer: {project.buyer or 'Not specified'}")
    document.add_paragraph(f"Deadline: {project.deadline or 'Not specified'}")
    document.add_paragraph(f"Approved items: {len(ready)}")
    document.add_paragraph(
        "Only responses individually approved by the workspace owner and "
        "marked Met are included. Recheck all current requirements and sources."
    )
    for index, (requirement, answer) in enumerate(ready, 1):
        document.add_heading(f"{index}. {requirement.category} requirement", level=2)
        document.add_paragraph(f"Source location: {requirement.locator}")
        document.add_paragraph("Requirement: " + requirement.evidence)
        document.add_paragraph("Approved response:")
        document.add_paragraph(answer.body)
        document.add_paragraph(f"Reviewer ID: {answer.approved_by or 'Unknown'}")
    output = BytesIO()
    document.save(output)
    return Response(
        output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": 'attachment; filename="trumonit-review-pack.docx"',
            "X-Content-Type-Options": "nosniff",
            "Cache-Control": "no-store",
        },
    )
