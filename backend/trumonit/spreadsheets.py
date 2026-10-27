"""Mapped RFP/RFQ XLSX imports and review-safe exports using original workbooks.

Questions and answer-cell locations are explicitly selected by the customer.
The product never invents answers, changes source questions or assumes a bid
is legally complete. Unsupported Excel features require manual verification.
"""
from __future__ import annotations

import hashlib
import io
import re
import zipfile
from datetime import date

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import Response
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .db import get_db
from .extract import MAX_DOCUMENT_BYTES, category
from .models import (
    AnswerLibraryEntry, Document, Requirement, RequirementResponse,
    SpreadsheetQuestion,
)
from .projects import owned_project, record
from .security import Principal, principal, require_editor

router = APIRouter(prefix="/projects", tags=["RFP spreadsheet response workflow"])

MAX_WORKBOOK_UNCOMPRESSED = 40 * 1024 * 1024
MAX_ZIP_MEMBER = 16 * 1024 * 1024
MAX_ZIP_MEMBERS = 500
MAX_WORKSHEETS = 20
MAX_SHEET_ROWS = 5000
MAX_SHEET_COLUMNS = 52
MAX_QUESTION_ROWS = 250


def validate_archive(data: bytes):
    if not data or len(data) > MAX_DOCUMENT_BYTES:
        raise HTTPException(413, "Maximum XLSX size is 8 MiB")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if not members or len(members) > MAX_ZIP_MEMBERS:
                raise HTTPException(422, "XLSX contains too many package parts")
            total = 0
            names = set()
            for part in members:
                name = part.filename.replace("\\", "/")
                if (name in names or name.startswith("/") or
                        any(segment in (".", "..") for segment in name.split("/"))):
                    raise HTTPException(422, "XLSX archive paths are unsafe")
                names.add(name)
                if part.flag_bits & 1:
                    raise HTTPException(422, "Encrypted XLSX files are not supported")
                if part.file_size > MAX_ZIP_MEMBER:
                    raise HTTPException(422, "XLSX archive member too large")
                if (part.file_size > 1024 * 1024 and
                        part.file_size > 200 * max(1, part.compress_size)):
                    raise HTTPException(422, "XLSX compression ratio exceeds the limit")
                total += part.file_size
                if total > MAX_WORKBOOK_UNCOMPRESSED:
                    raise HTTPException(422, "XLSX expanded size exceeds the limit")
            if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
                raise HTTPException(422, "Invalid XLSX workbook structure")
            if any(name.startswith((
                "xl/externalLinks/", "xl/embeddings/", "xl/activeX/",
                "xl/pivotTables/", "xl/connections",
            )) or name.endswith(("vbaProject.bin", ".vml")) for name in names):
                raise HTTPException(422, "Workbook contains unsupported linked or embedded features")
    except HTTPException:
        raise
    except (zipfile.BadZipFile, ValueError, OSError) as exc:
        raise HTTPException(422, "Invalid XLSX archive") from exc


def read_xlsx(data: bytes):
    validate_archive(data)
    try:
        book = load_workbook(io.BytesIO(data), read_only=False, data_only=False, keep_links=False)
        if len(book.worksheets) > MAX_WORKSHEETS:
            raise HTTPException(422, "Maximum 20 worksheets per workbook")
        for sheet in book.worksheets:
            if sheet.max_row > MAX_SHEET_ROWS or sheet.max_column > MAX_SHEET_COLUMNS:
                raise HTTPException(422, "Maximum sheet dimensions: 5000 rows and 52 columns")
        return book
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "Could not read standard XLSX file") from exc


def resolve_column(raw: str) -> tuple[str, int]:
    column = raw.strip().upper()
    if not re.fullmatch(r"[A-Z]{1,2}", column):
        raise HTTPException(422, "Column must be A through AZ")
    index = column_index_from_string(column)
    if index > MAX_SHEET_COLUMNS:
        raise HTTPException(422, "Column must be A through AZ")
    return column, index


async def file_bytes(file: UploadFile):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(415, "Only standard .xlsx workbooks are supported")
    if len(file.filename) > 180 or "\x00" in file.filename:
        raise HTTPException(422, "Invalid workbook filename")
    content = await file.read(MAX_DOCUMENT_BYTES + 1)
    if len(content) > MAX_DOCUMENT_BYTES or not content:
        raise HTTPException(413, "Maximum XLSX size is 8 MiB")
    return content


def inspect_workbook(book):
    sheets = []
    for sheet in book.worksheets:
        sample = []
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 12),
                                   max_col=min(sheet.max_column, 12)):
            values = [{"column": cell.column_letter, "value": str(cell.value)[:180]}
                      for cell in row if cell.value is not None]
            if values:
                sample.append({"row": row[0].row, "cells": values})
        sheets.append({"name": sheet.title, "max_row": sheet.max_row,
                       "max_column": sheet.max_column, "sample": sample,
                       "protected": bool(sheet.protection.sheet)})
    return sheets


@router.post("/{project_id}/spreadsheets/preview")
async def preview_spreadsheet(
    project_id: str,
    file: UploadFile = File(...),
    user: Principal = Depends(principal),
    db: Session = Depends(get_db),
):
    owned_project(db, project_id, user)
    require_editor(user)
    content = await file_bytes(file)
    book = read_xlsx(content)
    try:
        return {"sheets": inspect_workbook(book), "maximum_questions": MAX_QUESTION_ROWS,
                "notice": "Select the question and empty answer columns. Unsupported Excel features may not survive export."}
    finally:
        book.close()


@router.post("/{project_id}/spreadsheets", status_code=201)
async def import_spreadsheet(
    project_id: str,
    file: UploadFile = File(...),
    sheet_name: str = Form(..., min_length=1, max_length=31),
    question_column: str = Form(..., min_length=1, max_length=2),
    answer_column: str = Form(..., min_length=1, max_length=2),
    start_row: int = Form(..., ge=1, le=MAX_SHEET_ROWS),
    end_row: int = Form(..., ge=1, le=MAX_SHEET_ROWS),
    user: Principal = Depends(principal),
    db: Session = Depends(get_db),
):
    project = owned_project(db, project_id, user)
    require_editor(user)
    question_letter, qindex = resolve_column(question_column)
    answer_letter, aindex = resolve_column(answer_column)
    if qindex == aindex:
        raise HTTPException(422, "Question and answer columns must differ")
    if end_row < start_row or end_row - start_row + 1 > MAX_QUESTION_ROWS:
        raise HTTPException(422, "Choose a range of up to 250 rows")
    content = await file_bytes(file)
    digest = hashlib.sha256(content).hexdigest()
    if db.scalar(select(Document.id).where(
        Document.workspace_id == user.workspace_id,
        Document.project_id == project.id,
        Document.sha256 == digest,
    )):
        raise HTTPException(409, "Workbook was already imported into this project")
    count = db.scalar(select(func.count(Document.id)).where(
        Document.workspace_id == user.workspace_id,
        Document.project_id == project.id,
    )) or 0
    if count >= 25:
        raise HTTPException(422, "Maximum 25 documents per project")

    book = read_xlsx(content)
    try:
        if sheet_name not in book.sheetnames:
            raise HTTPException(422, "Selected worksheet does not exist")
        sheet = book[sheet_name]
        if sheet.protection.sheet:
            raise HTTPException(422, "Protected worksheets cannot be mapped")
        if end_row > sheet.max_row:
            raise HTTPException(422, "Selected range exceeds last populated worksheet row")
        mapped = []
        skipped_blank = 0
        for row_number in range(start_row, end_row + 1):
            question_cell = sheet.cell(row_number, qindex)
            raw = question_cell.value
            if raw is None or not str(raw).strip():
                skipped_blank += 1
                continue
            if question_cell.data_type == "f":
                raise HTTPException(422, f"Question row {row_number} contains a formula")
            question = " ".join(str(raw).split())
            if len(question) < 5 or len(question) > 700:
                raise HTTPException(422, f"Question row {row_number} must contain 5 to 700 characters")
            target = sheet.cell(row_number, aindex)
            if target.__class__.__name__ == "MergedCell":
                raise HTTPException(422, f"Answer cell {answer_letter}{row_number} is merged")
            if target.value is not None:
                raise HTTPException(422, f"Answer cell {answer_letter}{row_number} is not blank")
            mapped.append((row_number, question))
        if not mapped:
            raise HTTPException(422, "No non-empty questions in the selected rows")
    finally:
        book.close()

    name = file.filename.replace("\\", "/").split("/")[-1]
    document = Document(
        workspace_id=user.workspace_id, project_id=project.id,
        filename=name, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        sha256=digest, byte_count=len(content), content=content,
    )
    db.add(document)
    db.flush()
    for row_number, question in mapped:
        requirement = Requirement(
            workspace_id=user.workspace_id, project_id=project.id,
            document_id=document.id,
            locator=f"{sheet_name}!{question_letter}{row_number}",
            evidence=question, category=category(question),
        )
        db.add(requirement)
        db.flush()
        db.add(SpreadsheetQuestion(
            workspace_id=user.workspace_id, project_id=project.id,
            document_id=document.id, requirement_id=requirement.id,
            sheet_name=sheet_name, row_number=row_number,
            question_column=question_letter, answer_column=answer_letter,
        ))
    record(db, user, "spreadsheet.mapped_imported", document.id)
    db.commit()
    return {"document_id": document.id, "filename": document.filename,
            "questions_imported": len(mapped), "blank_rows_skipped": skipped_blank,
            "sheet_name": sheet_name,
            "note": "Questions were imported without AI. Review every question and approve responses before final export."}


def owned_spreadsheet(db: Session, project_id: str, document_id: str, user: Principal):
    owned_project(db, project_id, user)
    document = db.scalar(select(Document).where(
        Document.id == document_id, Document.project_id == project_id,
        Document.workspace_id == user.workspace_id,
        Document.content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ))
    if document is None:
        raise HTTPException(404, "Mapped spreadsheet not found")
    mappings = db.scalars(select(SpreadsheetQuestion).where(
        SpreadsheetQuestion.project_id == project_id,
        SpreadsheetQuestion.workspace_id == user.workspace_id,
        SpreadsheetQuestion.document_id == document_id,
    ).order_by(SpreadsheetQuestion.row_number)).all()
    if not mappings:
        raise HTTPException(404, "Mapped spreadsheet not found")
    return document, mappings


def exportable_answers(db: Session, user: Principal, project_id: str, mappings):
    reqids = [item.requirement_id for item in mappings]
    requirements = {item.id: item for item in db.scalars(select(Requirement).where(
        Requirement.id.in_(reqids), Requirement.workspace_id == user.workspace_id,
        Requirement.project_id == project_id,
    )).all()}
    responses = {item.requirement_id: item for item in db.scalars(select(RequirementResponse).where(
        RequirementResponse.requirement_id.in_(reqids),
        RequirementResponse.workspace_id == user.workspace_id,
        RequirementResponse.project_id == project_id,
    )).all()}
    linked = {row.library_entry_id for row in responses.values() if row.library_entry_id}
    library = {item.id: item for item in db.scalars(select(AnswerLibraryEntry).where(
        AnswerLibraryEntry.id.in_(linked), AnswerLibraryEntry.workspace_id == user.workspace_id,
    )).all()}
    today = date.today().isoformat()
    ready = {}
    for mapping in mappings:
        requirement = requirements.get(mapping.requirement_id)
        response = responses.get(mapping.requirement_id)
        if requirement is None or requirement.review_status != "met" or not response or (
            response.status != "approved" or not response.body.strip()
        ):
            continue
        if response.library_entry_id:
            source = library.get(response.library_entry_id)
            if (source is None or source.status != "approved" or (
                source.valid_until and source.valid_until < today
            )):
                continue
        ready[mapping.requirement_id] = response.body
    return ready


@router.get("/{project_id}/spreadsheets")
def list_spreadsheets(
    project_id: str, user: Principal = Depends(principal),
    db: Session = Depends(get_db),
):
    owned_project(db, project_id, user)
    mappings = db.scalars(select(SpreadsheetQuestion).where(
        SpreadsheetQuestion.workspace_id == user.workspace_id,
        SpreadsheetQuestion.project_id == project_id,
    ).order_by(SpreadsheetQuestion.document_id, SpreadsheetQuestion.row_number)).all()
    by_document = {}
    for row in mappings:
        by_document.setdefault(row.document_id, []).append(row)
    docs = {item.id: item for item in db.scalars(select(Document).where(
        Document.workspace_id == user.workspace_id, Document.project_id == project_id,
        Document.id.in_(by_document.keys()),
    )).all()}
    result = []
    for doc_id, entries in by_document.items():
        doc = docs.get(doc_id)
        if doc is None:
            continue
        ready = exportable_answers(db, user, project_id, entries)
        result.append({
            "document_id": doc_id, "filename": doc.filename,
            "sheet_name": entries[0].sheet_name,
            "question_count": len(entries), "approved_count": len(ready),
            "final_export_ready": len(ready) == len(entries),
        })
    return result


@router.get("/{project_id}/spreadsheets/{document_id}/export")
def export_spreadsheet(
    project_id: str, document_id: str,
    mode: str = Query("approved", pattern=r"^(approved|working)$"),
    user: Principal = Depends(principal),
    db: Session = Depends(get_db),
):
    document, mappings = owned_spreadsheet(db, project_id, document_id, user)
    ready = exportable_answers(db, user, project_id, mappings)
    if mode == "approved" and len(ready) != len(mappings):
        raise HTTPException(409, f"Workbook requires {len(mappings) - len(ready)} more reviewed and approved responses")
    book = read_xlsx(document.content)
    try:
        for mapping in mappings:
            text = ready.get(mapping.requirement_id)
            if text is None:
                continue
            if mapping.sheet_name not in book.sheetnames:
                raise HTTPException(409, "Original sheet is missing")
            destination = book[mapping.sheet_name][f"{mapping.answer_column}{mapping.row_number}"]
            if destination.value is not None or destination.__class__.__name__ == "MergedCell":
                raise HTTPException(409, "Original answer destination is no longer blank")
            destination.value = text
            # Prevent an answer that begins with = from being interpreted as
            # an Excel formula. Plain strings remain plain strings.
            destination.data_type = "s"
        stream = io.BytesIO()
        book.save(stream)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "Cannot safely export this workbook") from exc
    finally:
        book.close()
    safe_filename = re.sub(r"[^A-Za-z0-9_.-]", "_", document.filename.rsplit(".", 1)[0])
    label = "approved" if mode == "approved" else "working"
    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{safe_filename}-{label}.xlsx"',
            "X-Content-Type-Options": "nosniff",
            "Cache-Control": "no-store",
            "X-Trumonit-Export-Status": label,
        },
    )
