import io
import zipfile

import fitz
import pytest
from docx import Document as WordDocument
from fastapi import HTTPException

from trumonit.extract import candidate_segments


def test_pdf_evidence_preserves_page_number_and_does_not_make_claims():
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((65, 85), "The applicant must supply a current insurance certificate.")
    page.insert_text((65, 115), "Reference material is provided for information.")
    output = doc.tobytes()
    doc.close()
    candidates = candidate_segments("rfp.pdf", output)
    assert len(candidates) == 1
    assert candidates[0].locator == "page 1"
    assert candidates[0].category == "legal"
    assert "must supply" in candidates[0].evidence


def test_docx_and_unsupported_input():
    doc = WordDocument()
    doc.add_paragraph("Wykonawca musi przedstawić ważne zaświadczenie.")
    doc.add_paragraph("This line is general context.")
    buf = io.BytesIO()
    doc.save(buf)
    found = candidate_segments("specification.docx", buf.getvalue())
    assert len(found) == 1
    assert found[0].locator.startswith("paragraph")
    assert found[0].category == "legal"
    with pytest.raises(HTTPException) as exc:
        candidate_segments("payload.exe", b"not-a-document")
    assert exc.value.status_code == 415


def test_docx_rejects_zip_bomb_before_xml_expansion():
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", compression=zipfile.ZIP_DEFLATED) as stream:
        stream.writestr("[Content_Types].xml", "<Types/>")
        stream.writestr("word/document.xml", "A" * (17 * 1024 * 1024))
    assert len(archive.getvalue()) < 8 * 1024 * 1024
    with pytest.raises(HTTPException) as exc:
        candidate_segments("compressed.docx", archive.getvalue())
    assert exc.value.status_code == 422


def test_docx_rejects_duplicate_members_and_path_traversal():
    import warnings
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as stream:
        stream.writestr("[Content_Types].xml", "<Types/>")
        stream.writestr("word/document.xml", "<document/>")
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            stream.writestr("word/document.xml", "<document/>")
    with pytest.raises(HTTPException) as exc:
        candidate_segments("duplicate.docx", archive.getvalue())
    assert exc.value.status_code == 422
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w") as stream:
        stream.writestr("[Content_Types].xml", "<Types/>")
        stream.writestr("word/document.xml", "<document/>")
        stream.writestr("../unexpected.xml", "<danger/>")
    with pytest.raises(HTTPException) as exc:
        candidate_segments("traversal.docx", archive.getvalue())
    assert exc.value.status_code == 422
