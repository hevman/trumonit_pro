import io

from docx import Document as WordDocument
from .conftest import authorized, register


def docx_bytes():
    doc = WordDocument()
    doc.add_paragraph("The supplier must provide proof of current professional insurance.")
    doc.add_paragraph("The final proposal deadline is 31 December 2027.")
    doc.add_paragraph("Informational paragraph without an obligation.")
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def test_register_login_project_upload_review_and_export(client, gateway):
    first = register(client, gateway)
    token = first["session_token"]
    owner = authorized(gateway, token)
    me = client.get("/api/auth/me", headers=owner)
    assert me.status_code == 200
    assert me.json()["workspace_name"] == "Example Company"

    created = client.post("/api/projects", headers=owner, json={
        "name": "Industrial service RFP",
        "buyer": "Example public purchaser",
        "deadline": "2027-12-31",
    })
    assert created.status_code == 201, created.text
    project = created.json()
    assert project["stage"] == "review"
    project_id = project["id"]

    content = docx_bytes()
    uploaded = client.post(
        f"/api/projects/{project_id}/documents",
        headers=owner,
        files={"file": ("terms.docx", content,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert uploaded.status_code == 201, uploaded.text
    assert uploaded.json()["requirements_created"] >= 2
    assert "unverified candidates" in uploaded.json()["note"]

    duplicate = client.post(
        f"/api/projects/{project_id}/documents",
        headers=owner, files={"file": ("same.docx", content)},
    )
    assert duplicate.status_code == 409

    detail = client.get(f"/api/projects/{project_id}", headers=owner)
    assert detail.status_code == 200
    assert detail.json()["requirements_total"] >= 2
    requirement = detail.json()["requirements"][0]
    assert requirement["locator"].startswith("paragraph ")

    changed = client.patch(
        f"/api/projects/{project_id}/requirements/{requirement['id']}",
        headers=owner,
        json={"status": "met", "note": "Certificate attached to submission."},
    )
    assert changed.status_code == 200, changed.text
    assert changed.json()["status"] == "met"

    export = client.get(f"/api/projects/{project_id}/export", headers=owner)
    assert export.status_code == 200
    assert "source,evidence,status" in export.text
    assert "Certificate attached" in export.text

    download = client.get(
        f"/api/projects/{project_id}/documents/{uploaded.json()['id']}",
        headers=owner,
    )
    assert download.status_code == 200
    assert download.content == content
    assert download.headers["x-content-type-options"] == "nosniff"

    logout = client.post("/api/auth/logout", headers=owner)
    assert logout.status_code == 200
    assert client.get("/api/auth/me", headers=owner).status_code == 401


def test_tenant_isolation_and_missing_gateway(client, gateway):
    a = register(client, gateway, "first@example.test", "First Company")
    b = register(client, gateway, "second@example.test", "Second Company")
    a_headers = authorized(gateway, a["session_token"])
    b_headers = authorized(gateway, b["session_token"])
    created = client.post("/api/projects", headers=a_headers,
                          json={"name": "Private project"}).json()
    pid = created["id"]
    assert len(client.get("/api/projects", headers=a_headers).json()) == 1
    assert len(client.get("/api/projects", headers=b_headers).json()) == 0
    assert client.get(f"/api/projects/{pid}", headers=b_headers).status_code == 404
    assert client.patch(f"/api/projects/{pid}", headers=b_headers,
                        json={"stage": "submitted"}).status_code == 404
    assert client.post(f"/api/projects/{pid}/documents", headers=b_headers,
                       files={"file": ("secret.docx", docx_bytes())}).status_code == 404
    assert client.get(f"/api/projects/{pid}/export", headers=b_headers).status_code == 404
    assert client.get("/api/projects").status_code == 401
    assert client.get("/api/projects", headers=gateway).status_code == 401
    assert client.get("/health").status_code == 200


def test_login_rejects_invalid_password_and_duplicate_registration(client, gateway):
    first = register(client, gateway)
    duplicate = client.post("/api/auth/register", headers=gateway, json={
        "name": "Other user", "email": "OWNER@example.test",
        "password": "ValidPassword2026!", "workspace_name": "Duplicate company",
    })
    assert duplicate.status_code == 409
    invalid = client.post("/api/auth/login", headers=gateway,
                          json={"email": "owner@example.test", "password": "invalid"})
    assert invalid.status_code == 401
    valid = client.post("/api/auth/login", headers=gateway, json={
        "email": "owner@example.test", "password": "VeryStrongPassword2026!",
    })
    assert valid.status_code == 200
    assert valid.json()["session_token"] != first["session_token"]
