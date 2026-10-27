"""Customer XLSX RFP: original workbook -> team answers -> approved Excel."""
import io
import zipfile

import openpyxl
from openpyxl.styles import Font, PatternFill

from .conftest import authorized, register


MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook_bytes():
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Security Questions"
    sheet["A1"] = "ID"
    sheet["B1"] = "Question"
    sheet["C1"] = "Supplier response"
    sheet["D1"] = "Scoring"
    sheet["A2"] = 1
    sheet["B2"] = "Does your company hold current insurance documentation?"
    sheet["A3"] = 2
    sheet["B3"] = "Describe your organization's security management policy."
    sheet["D2"] = "=1+1"
    sheet["B2"].font = Font(bold=True)
    sheet["C2"].fill = PatternFill(fill_type="solid", fgColor="FFDDEEFF")
    out = io.BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()


def prepared(client, gateway, email="xlsx-owner@example.test"):
    info = register(client, gateway, email, "International supplier")
    headers = authorized(gateway, info["session_token"])
    project = client.post("/api/projects", headers=headers,
                          json={"name": "International security RFP"})
    assert project.status_code == 201, project.text
    return headers, project.json()["id"]


def import_file(client, owner, project, content=None, **mapping):
    return client.post(f"/api/projects/{project}/spreadsheets", headers=owner,
                       data={
                           "sheet_name": "Security Questions",
                           "question_column": "B", "answer_column": "C",
                           "start_row": "2", "end_row": "3", **mapping,
                       },
                       files={"file": ("customer-rfp.xlsx", content if content is not None else workbook_bytes(), MIME)})


def test_xlsx_preserves_layout_and_needs_each_owner_approval(client, gateway):
    owner, project = prepared(client, gateway)
    raw = workbook_bytes()
    preview = client.post(f"/api/projects/{project}/spreadsheets/preview",
                          headers=owner, files={"file": ("customer-rfp.xlsx", raw, MIME)})
    assert preview.status_code == 200, preview.text
    assert preview.json()["sheets"][0]["name"] == "Security Questions"
    assert preview.json()["sheets"][0]["sample"][1]["cells"][1]["column"] == "B"

    uploaded = import_file(client, owner, project, raw)
    assert uploaded.status_code == 201, uploaded.text
    assert uploaded.json()["questions_imported"] == 2
    document_id = uploaded.json()["document_id"]
    assert import_file(client, owner, project, raw).status_code == 409

    detail = client.get(f"/api/projects/{project}", headers=owner)
    assert detail.status_code == 200
    rows = detail.json()["requirements"]
    assert len(rows) == 2
    assert rows[0]["locator"] == "Security Questions!B2"
    assert rows[1]["locator"] == "Security Questions!B3"

    listing = client.get(f"/api/projects/{project}/spreadsheets", headers=owner)
    assert listing.status_code == 200
    assert listing.json()[0]["question_count"] == 2
    assert listing.json()[0]["approved_count"] == 0
    assert listing.json()[0]["final_export_ready"] is False

    final_url = f"/api/projects/{project}/spreadsheets/{document_id}/export"
    assert client.get(final_url, headers=owner).status_code == 409
    working = client.get(final_url+"?mode=working", headers=owner)
    assert working.status_code == 200
    book = openpyxl.load_workbook(io.BytesIO(working.content))
    assert book["Security Questions"]["C2"].value is None
    assert book["Security Questions"]["D2"].value == "=1+1"
    book.close()

    for index, row in enumerate(rows):
        rid = row["id"]
        assert client.patch(
            f"/api/projects/{project}/requirements/{rid}", headers=owner,
            json={"status": "met", "note": "Verified by reviewer"},
        ).status_code == 200
        answer = '=HYPERLINK("https://example.org","literal")' if index == 0 else (
            "Our organization maintains a documented policy and review process."
        )
        draft = client.patch(
            f"/api/projects/{project}/requirements/{rid}/response",
            headers=owner, json={"body": answer},
        )
        assert draft.status_code == 200, draft.text
        assert client.get(final_url, headers=owner).status_code == 409
        approved = client.post(
            f"/api/projects/{project}/requirements/{rid}/response/approve",
            headers=owner,
        )
        assert approved.status_code == 200, approved.text

    result = client.get(final_url, headers=owner)
    assert result.status_code == 200, result.text
    assert result.headers["x-trumonit-export-status"] == "approved"
    returned = openpyxl.load_workbook(io.BytesIO(result.content))
    sheet = returned["Security Questions"]
    assert sheet["B2"].value == "Does your company hold current insurance documentation?"
    assert sheet["C2"].value == '=HYPERLINK("https://example.org","literal")'
    assert sheet["C2"].data_type == "s"
    assert sheet["C2"].fill.fgColor.rgb == "FFDDEEFF"
    assert sheet["B2"].font.bold is True
    assert sheet["D2"].value == "=1+1"
    assert sheet["C3"].value.startswith("Our organization")
    returned.close()
    assert client.get(f"/api/projects/{project}/spreadsheets", headers=owner).json()[0][
        "final_export_ready"] is True
    source = client.get(f"/api/projects/{project}/documents/{document_id}", headers=owner)
    assert source.status_code == 200
    assert source.content == raw


def test_workbook_mapping_rejects_mistakes_and_cross_tenant_access(client, gateway):
    first, project_a = prepared(client, gateway, "xlsx-first@example.test")
    second, project_b = prepared(client, gateway, "xlsx-second@example.test")
    original = workbook_bytes()
    assert client.post(f"/api/projects/{project_a}/spreadsheets/preview",
                       headers=second, files={"file": ("form.xlsx", original)}).status_code == 404
    assert import_file(client, first, project_a, question_column="C",
                       answer_column="C").status_code == 422
    assert import_file(client, first, project_a,
                       sheet_name="Other sheet").status_code == 422
    assert import_file(client, first, project_a, start_row="2",
                       end_row="300").status_code == 422
    assert import_file(client, first, project_a, question_column="A").status_code == 422

    book = openpyxl.load_workbook(io.BytesIO(original))
    book.active["C2"] = "Do not overwrite buyer content"
    out = io.BytesIO()
    book.save(out)
    book.close()
    assert import_file(client, first, project_a, out.getvalue()).status_code == 422

    imported = import_file(client, first, project_a)
    assert imported.status_code == 201, imported.text
    document_id = imported.json()["document_id"]
    assert client.get(f"/api/projects/{project_a}/spreadsheets", headers=second).status_code == 404
    assert client.get(f"/api/projects/{project_a}/spreadsheets/{document_id}/export?mode=working",
                      headers=second).status_code == 404
    assert client.get(f"/api/projects/{project_b}/spreadsheets/{document_id}/export?mode=working",
                      headers=second).status_code == 404
    deleted = client.post(f"/api/projects/{project_a}/delete", headers=first, json={
        "confirm_project_name": "International security RFP",
    })
    assert deleted.status_code == 200, deleted.text
    assert client.get(f"/api/projects/{project_a}/spreadsheets", headers=first).status_code == 404


def test_xlsx_decompression_bomb_rejected(client, gateway):
    owner, project = prepared(client, gateway)
    binary = io.BytesIO()
    with zipfile.ZipFile(binary, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("xl/workbook.xml", "B" * (17 * 1024 * 1024))
    assert len(binary.getvalue()) < 8 * 1024 * 1024
    result = client.post(f"/api/projects/{project}/spreadsheets/preview",
                         headers=owner, files={"file": ("untrusted.xlsx", binary.getvalue())})
    assert result.status_code == 422


def test_library_approval_revocation_blocks_final_xlsx(client, gateway):
    owner, project = prepared(client, gateway)
    document_id = import_file(client, owner, project).json()["document_id"]
    rows = client.get(f"/api/projects/{project}", headers=owner).json()["requirements"]
    library = client.post("/api/library", headers=owner, json={
        "title": "Reusable supplier assurance statement",
        "body": "We maintain a documented insurance policy for our client contracts.",
    }).json()
    assert client.post(f"/api/library/{library['id']}/approve", headers=owner).status_code == 200
    for row in rows:
        rid = row["id"]
        assert client.patch(f"/api/projects/{project}/requirements/{rid}",
                            headers=owner, json={"status": "met", "note": "Reviewed"}).status_code == 200
        assert client.patch(f"/api/projects/{project}/requirements/{rid}/response",
                            headers=owner, json={"library_entry_id": library["id"]}).status_code == 200
        assert client.post(f"/api/projects/{project}/requirements/{rid}/response/approve",
                           headers=owner).status_code == 200
    url = f"/api/projects/{project}/spreadsheets/{document_id}/export"
    assert client.get(url, headers=owner).status_code == 200
    edited = client.patch(f"/api/library/{library['id']}", headers=owner, json={
        "title": "Reusable supplier assurance statement",
        "body": "Our evidence has changed and requires another review.",
    })
    assert edited.status_code == 200, edited.text
    assert client.get(url, headers=owner).status_code == 409
    working = client.get(url+"?mode=working", headers=owner)
    assert working.status_code == 200
    book = openpyxl.load_workbook(io.BytesIO(working.content))
    assert book.active["C2"].value is None
    book.close()
